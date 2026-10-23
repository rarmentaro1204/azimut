#!/usr/bin/env python3
"""Scraper catalogo Azimut Yachts -> Excel (foglio Macchine) + immagini + link alle brochure.

Stesso schema colonne del catalogo Ferretti (excel/catalogo_*.xlsx), così tools/build_catalog_js.py
e il configuratore lavorano senza modifiche. Solo libreria standard + openpyxl.
Pausa tra richieste, rispetta robots.txt, nessun login. Prezzo = 0 (da compilare con il listino).

Strategia (il sito non è stato ispezionato quando lo script è stato scritto: verificare con --dry-run):
  1. sitemap (robots.txt -> Sitemap:, oppure /sitemap.xml) -> URL che sembrano pagine modello;
  2. per ogni pagina: JSON-LD schema.org (Product/Vehicle) se presente, altrimenti meta OpenGraph;
  3. specifiche tecniche cercate per etichetta nel testo (LOA, baglio, pescaggio, dislocamento, serbatoi, cabine, motori);
  4. link a PDF "brochure"/"catalogue"/"scheda tecnica" raccolti nella colonna Brochure_URL e scaricati con --brochures.
Modalità offline: --html-dir DIR legge file .html salvati a mano (nome file = codice modello) invece di scaricare.
Uso: python3 scrape_azimut.py [--base https://www.azimutyachts.com] [--lang it-it] [--out FILE.xlsx]
                             [--delay 1.5] [--images DIR] [--brochures DIR] [--dry-run] [--html-dir DIR]
"""
import argparse, html as H, json, os, re, sys, time, urllib.parse, urllib.request, urllib.robotparser
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

UA = "Mozilla/5.0 (compatible; CatalogBot/1.0)"
COLONNE = ["Codice", "Famiglia", "Categoria", "Sottocategoria", "Nome", "Prezzo",
           "LOA_m", "Larghezza_m", "Pescaggio_m", "Lunghezza_scafo_m", "Cabine",
           "Motore_Standard", "Vel_Max_Standard_kn", "Vel_Crociera_Standard_kn",
           "Motore_Optional", "Vel_Max_Optional_kn", "Vel_Crociera_Optional_kn",
           "Dislocamento_Vuoto_kg", "Dislocamento_Carico_kg", "Carburante_l", "Acqua_l", "Persone",
           "Descrizione", "Descrizione_EN", "Storia_IT", "Storia_EN", "URL", "Immagine_URL", "Immagine_File",
           "Brochure_URL"]

# etichette (it/en) -> colonna; il valore è la riga/cella successiva o dopo ':'
ETICHETTE = {
    "LOA_m": ("lunghezza fuori tutto", "length overall", "loa", "lunghezza f.t."),
    "Lunghezza_scafo_m": ("lunghezza scafo", "hull length", "lh"),
    "Larghezza_m": ("larghezza massima", "baglio massimo", "baglio", "beam", "max beam"),
    "Pescaggio_m": ("pescaggio", "draft", "draught"),
    "Dislocamento_Vuoto_kg": ("dislocamento a vuoto", "light displacement", "dry weight", "peso a vuoto"),
    "Dislocamento_Carico_kg": ("dislocamento a pieno carico", "full load displacement", "dislocamento"),
    "Carburante_l": ("serbatoi carburante", "serbatoio carburante", "fuel tank", "fuel capacity", "carburante"),
    "Acqua_l": ("serbatoi acqua", "serbatoio acqua", "water tank", "fresh water", "acqua"),
    "Persone": ("persone a bordo", "numero persone", "max persons", "persons on board", "persone"),
    "Cabine": ("cabine", "cabins"),
}
SERIE_NOTE = ("flybridge", "magellano", "grande", "atlantis", "s", "seadeck", "verve", "fly", "azimut")
PAROLE_BROCHURE = ("brochure", "catalog", "scheda", "datasheet", "specification", "depliant")


def get(url, binary=False):
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    for t in range(3):
        try:
            with urllib.request.urlopen(req, timeout=40) as r:
                d = r.read()
                return d if binary else d.decode("utf-8", "ignore")
        except Exception:
            if t == 2:
                raise
            time.sleep(2 * (t + 1))


def num(s):
    m = re.search(r"-?\d+(?:[.,]\d+)?", str(s or "").replace("\xa0", " "))
    if not m:
        return None
    t = m.group(0)
    if re.fullmatch(r"\d{1,3}[.]\d{3}", t):          # 1.250 = migliaia
        t = t.replace(".", "")
    return float(t.replace(",", "."))


def righe(page):
    page = re.sub(r"<script.*?</script>|<style.*?</style>|<noscript.*?</noscript>", "", page, flags=re.S | re.I)
    page = re.sub(r"</(td|th|li|p|div|tr|h\d)>", "\n", page, flags=re.I)
    return [re.sub(r"\s+", " ", H.unescape(l)).strip() for l in re.sub(r"<[^>]+>", "\n", page).split("\n") if l.strip()]


def valore_dopo(L, etichette, start=0):
    """Valore associato a un'etichetta: stessa riga dopo ':' oppure riga successiva."""
    for i in range(start, len(L)):
        low = L[i].lower().rstrip(":")
        for e in etichette:
            if low == e:
                return L[i + 1] if i + 1 < len(L) else ""
            if low.startswith(e + ":") or low.startswith(e + " :"):
                return L[i].split(":", 1)[1].strip()
    return ""


def jsonld(page):
    out = []
    for m in re.finditer(r'<script[^>]+type="application/ld\+json"[^>]*>(.*?)</script>', page, flags=re.S | re.I):
        try:
            d = json.loads(H.unescape(m.group(1)).strip())
        except Exception:
            continue
        stack = d if isinstance(d, list) else [d]
        while stack:
            x = stack.pop()
            if isinstance(x, dict):
                out.append(x)
                stack.extend(v for v in x.values() if isinstance(v, (dict, list)))
            elif isinstance(x, list):
                stack.extend(x)
    return out


def meta(page, prop):
    m = re.search(r'<meta[^>]+(?:property|name)="%s"[^>]+content="([^"]*)"' % re.escape(prop), page, flags=re.I) \
        or re.search(r'<meta[^>]+content="([^"]*)"[^>]+(?:property|name)="%s"' % re.escape(prop), page, flags=re.I)
    return H.unescape(m.group(1)).strip() if m else ""


def robots(base):
    rp = urllib.robotparser.RobotFileParser()
    try:
        rp.parse(get(base + "/robots.txt").splitlines())
    except Exception:
        rp.parse([])
    return rp


def sitemap_urls(base, rp):
    mappe = re.findall(r"(?im)^sitemap:\s*(\S+)", get(base + "/robots.txt")) or [base + "/sitemap.xml"]
    visti, urls = set(), []
    while mappe:
        m = mappe.pop(0)
        if m in visti:
            continue
        visti.add(m)
        try:
            xml = get(m)
        except Exception as e:
            print("sitemap KO", m, e, file=sys.stderr)
            continue
        for loc in re.findall(r"<loc>\s*([^<\s]+)\s*</loc>", xml):
            (mappe if loc.endswith(".xml") else urls).append(loc)
        time.sleep(0.5)
    return [u for u in dict.fromkeys(urls) if rp.can_fetch(UA, u)]


def sembra_modello(url, lang):
    """Euristica: percorso con lingua e un segmento che contiene una cifra (es. /yachts/flybridge/azimut-60)."""
    p = urllib.parse.urlparse(url).path.lower()
    if lang and lang.lower() not in p:
        return False
    return bool(re.search(r"/(yachts?|models?|modelli|collection|collezion\w*|imbarcazion\w*)/", p)) and bool(re.search(r"\d", p.rsplit("/", 1)[-1] or p))


def codice_da(url, nome):
    slug = urllib.parse.urlparse(url).path.rstrip("/").rsplit("/", 1)[-1]
    return re.sub(r"[^A-Za-z0-9]+", "-", slug or nome).strip("-").upper()


def famiglia_da(url, nome):
    p = urllib.parse.urlparse(url).path.lower()
    for s in ("flybridge", "magellano", "grande", "atlantis", "s-collection", "seadeck", "verve", "fly"):
        if s in p or s in nome.lower():
            return s.replace("-collection", "").capitalize()
    return "Azimut"


def parse_modello(url, page, base):
    L = righe(page)
    ld = jsonld(page)
    prod = next((x for x in ld if str(x.get("@type", "")).lower() in ("product", "vehicle", "boat")), {})
    nome = H.unescape(str(prod.get("name") or meta(page, "og:title") or (L[0] if L else ""))).split("|")[0].strip()
    desc = H.unescape(str(prod.get("description") or meta(page, "og:description") or meta(page, "description")))
    img = prod.get("image") or meta(page, "og:image")
    if isinstance(img, list):
        img = img[0] if img else ""
    if isinstance(img, dict):
        img = img.get("url", "")
    r = {c: "" for c in COLONNE}
    fam = famiglia_da(url, nome)
    r.update(Codice=codice_da(url, nome), Famiglia=fam, Categoria="Motor yacht", Sottocategoria=fam,
             Nome=nome, Prezzo=0, Descrizione=desc, URL=url,
             Immagine_URL=urllib.parse.urljoin(base, img) if img else "")
    for col, ets in ETICHETTE.items():
        v = valore_dopo(L, ets)
        r[col] = (v if col == "Cabine" else num(v)) or ""
    pdfs = [urllib.parse.urljoin(url, h) for h in re.findall(r'href="([^"]+\.pdf[^"]*)"', page, flags=re.I)]
    br = [p for p in pdfs if any(k in p.lower() for k in PAROLE_BROCHURE)] or pdfs
    r["Brochure_URL"] = br[0] if br else ""
    # motori / velocità: righe che contengono marca+potenza e kn
    mot = next((l for l in L if re.search(r"\b(\d{3,4})\s*(hp|cv|kw)\b", l, re.I) and re.search(r"volvo|man|mtu|cat|cummins|yanmar|ips|d\d", l, re.I)), "")
    r["Motore_Standard"] = mot
    kn = [num(x) for x in re.findall(r"(\d+(?:[.,]\d+)?)\s*(?:kn|kt|nodi|knots)", " ".join(L), flags=re.I)]
    if kn:
        r["Vel_Max_Standard_kn"] = max(kn)
        r["Vel_Crociera_Standard_kn"] = sorted(kn)[len(kn) // 2] if len(kn) > 1 else ""
    r["Storia_IT"] = "\n\n".join(list(dict.fromkeys(l for l in L if len(l) > 120))[:4])
    return r


def scrivi_excel(rows, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Macchine"
    ws["A1"] = "Catalogo Azimut Yachts (estratto dal sito ufficiale)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Prezzi a 0: da compilare. Dati tecnici dal sito, da verificare con il listino/brochure ufficiale prima di emettere offerte."
    for c, nome in enumerate(COLONNE, 1):
        cell = ws.cell(row=4, column=c, value=nome)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C21632")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        lungo = nome in ("Descrizione", "Descrizione_EN", "Storia_IT", "Storia_EN", "URL", "Immagine_URL", "Brochure_URL")
        ws.column_dimensions[get_column_letter(c)].width = 40 if lungo else 22
    for ri, row in enumerate(rows, 5):
        for c, nome in enumerate(COLONNE, 1):
            ws.cell(row=ri, column=c, value=row.get(nome, ""))
    ws.freeze_panes = "F5"
    wb.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default="https://www.azimutyachts.com")
    ap.add_argument("--lang", default="", help="segmento lingua nell'URL, es. it-it (vuoto = tutti)")
    ap.add_argument("--out", default="../excel/catalogo_azimut.xlsx")
    ap.add_argument("--delay", type=float, default=1.5)
    ap.add_argument("--images", default="../images")
    ap.add_argument("--brochures", default="", help="se indicato, scarica i PDF trovati in questa cartella")
    ap.add_argument("--html-dir", default="", help="modalità offline: usa i .html salvati (nome file = codice)")
    ap.add_argument("--dry-run", action="store_true", help="elenca solo gli URL modello trovati")
    a = ap.parse_args()
    base = a.base.rstrip("/")
    pagine = []   # (url, html)
    if a.html_dir:
        for f in sorted(os.listdir(a.html_dir)):
            if f.endswith(".html"):
                pagine.append((f"{base}/{os.path.splitext(f)[0]}", open(os.path.join(a.html_dir, f), encoding="utf-8").read()))
    else:
        rp = robots(base)
        urls = [u for u in sitemap_urls(base, rp) if sembra_modello(u, a.lang)]
        print(f"{len(urls)} URL candidati")
        if a.dry_run:
            print("\n".join(urls))
            return
        for u in urls:
            try:
                pagine.append((u, get(u)))
            except Exception as e:
                print("ERRORE", u, e, file=sys.stderr)
            time.sleep(a.delay)
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    os.makedirs(a.images, exist_ok=True)
    rows, visti = [], set()
    for url, page in pagine:
        r = parse_modello(url, page, base)
        if not r["Nome"] or r["Codice"] in visti or not (r["LOA_m"] or r["Immagine_URL"]):
            continue                     # pagina non-modello (nessuna misura né immagine)
        visti.add(r["Codice"])
        if r["Immagine_URL"] and not a.html_dir:
            ext = os.path.splitext(urllib.parse.urlparse(r["Immagine_URL"]).path)[1] or ".jpg"
            f = f"{r['Codice']}{ext}"
            try:
                with open(os.path.join(a.images, f), "wb") as fh:
                    fh.write(get(r["Immagine_URL"], binary=True))
                r["Immagine_File"] = f
            except Exception as e:
                print("img KO", r["Codice"], e, file=sys.stderr)
        if a.brochures and r["Brochure_URL"] and not a.html_dir:
            os.makedirs(a.brochures, exist_ok=True)
            try:
                with open(os.path.join(a.brochures, f"{r['Codice']}.pdf"), "wb") as fh:
                    fh.write(get(r["Brochure_URL"], binary=True))
            except Exception as e:
                print("brochure KO", r["Codice"], e, file=sys.stderr)
        rows.append(r)
        print(r["Codice"], r["Nome"], r["LOA_m"], r["Cabine"])
    scrivi_excel(rows, a.out)
    print("scritto", a.out, len(rows))


if __name__ == "__main__":
    main()
